"""
resume_parser.py — Extract structured profile data from a résumé.

Two responsibilities:

  1. extract_text()  — pull plain text out of an uploaded PDF / DOCX / TXT.
  2. parse_resume()  — ask the local LLM to map that text onto the
     UserProfile schema, strictly EXTRACTIVELY (never invent data).

Used only by the first-run onboarding wizard. The result is shown to the
user for review/edit before anything is persisted — this module never
writes config.
"""

from __future__ import annotations

import io
import json
import logging
import re

from . import ai_engine
from . import apple_bridge
from . import prompt_registry

logger = logging.getLogger(__name__)

# Fields the wizard can prefill from a résumé. Demographic / credential /
# salary fields are intentionally excluded — they are not on a résumé and
# must be entered deliberately by the user.
_STR_FIELDS = (
    "full_name", "email", "phone", "location",
    "street_address", "street_address_2", "city", "state", "zip_code",
    "linkedin", "github", "portfolio", "summary",
)
_MAX_CHARS = 16000  # keep the prompt inside the local model's context window
# Apple's on-device model takes at most 8,000 characters of input
# (apple-bridge OnDeviceModel.swift). With it as the strong tier the résumé is
# parsed in section-aligned chunks under 7,000 characters (less the prompt) and
# the partial profiles are merged. The iOS twin is ResumeProfileParser.chunk.
APPLE_INPUT_CAP = 8000
APPLE_CHUNK_CHARS = 7000
_HEADINGS = {
    "summary", "professional summary", "profile", "objective", "experience", "work experience",
    "professional experience", "relevant experience", "employment", "employment history", "work history",
    "education", "skills", "technical skills", "core skills", "certifications", "certificates",
    "licenses", "licenses and certifications", "projects", "awards", "publications", "volunteer",
    "volunteer experience", "languages", "interests", "references",
}


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------
def extract_text(filename: str, data: bytes) -> str:
    """Return plain text from an uploaded résumé.

    Supports .pdf (pypdf), .docx (python-docx) and .txt / plain text.
    Raises ValueError for unsupported or unreadable files.
    """
    name = (filename or "").lower().strip()

    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise ValueError("PDF support requires the 'pypdf' package") from exc
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [(page.extract_text() or "") for page in reader.pages]
        except Exception as exc:
            raise ValueError(f"Could not read PDF: {exc}") from exc
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise ValueError("DOCX support requires the 'python-docx' package") from exc
        try:
            doc = Document(io.BytesIO(data))
        except Exception as exc:
            raise ValueError(f"Could not read DOCX: {exc}") from exc
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append("  ".join(cells))
        return "\n".join(lines).strip()

    if name.endswith((".txt", ".md", ".text")) or not name:
        try:
            return data.decode("utf-8", errors="replace").strip()
        except Exception as exc:
            raise ValueError(f"Could not decode text file: {exc}") from exc

    raise ValueError(
        f"Unsupported file type: {filename!r}. Upload a PDF, DOCX, or TXT, "
        "or paste the résumé text instead."
    )


# ---------------------------------------------------------------------------
# LLM extraction
# ---------------------------------------------------------------------------
# The extraction prompt lives in prompt_registry (key "resume_parse") so it
# can be edited from Settings → Prompts. Other extractive sources (e.g. the
# LinkedIn importer) pass their own registry key to parse_resume().


def _flatten_item(v) -> str:
    """One display string per item; models sometimes return objects
    (e.g. {"name": "Security+", "issuer": "CompTIA"}) despite the prompt."""
    if isinstance(v, dict):
        parts = [str(x).strip() for x in v.values()
                 if isinstance(x, (str, int, float)) and str(x).strip()]
        return " — ".join(parts)
    return str(v).strip()


def _coerce_str_list(value) -> list[str]:
    if isinstance(value, list):
        return [s for s in (_flatten_item(v) for v in value) if s]
    if isinstance(value, str) and value.strip():
        return [s.strip() for s in re.split(r"[,\n;]", value) if s.strip()]
    return []


def _sanitize(raw: dict) -> dict:
    """Coerce the model's JSON onto the UserProfile partial shape.

    Drops unknown keys, fixes types, and removes empty experience/education
    rows so the review screen isn't littered with blanks.
    """
    out: dict = {}
    for f in _STR_FIELDS:
        v = raw.get(f, "")
        out[f] = v.strip() if isinstance(v, str) else ("" if v is None else str(v))

    out["skills"] = _coerce_str_list(raw.get("skills"))
    out["certifications"] = _coerce_str_list(raw.get("certifications"))

    experience = []
    for e in raw.get("experience") or []:
        if not isinstance(e, dict):
            continue
        title = str(e.get("title", "")).strip()
        company = str(e.get("company", "")).strip()
        if not title and not company:
            continue
        experience.append({
            "title": title,
            "company": company,
            "start_date": str(e.get("start_date", "")).strip(),
            "end_date": str(e.get("end_date", "") or "Present").strip(),
            "bullets": _coerce_str_list(e.get("bullets")),
        })
    out["experience"] = experience

    education = []
    for e in raw.get("education") or []:
        if not isinstance(e, dict):
            continue
        degree = str(e.get("degree", "")).strip()
        school = str(e.get("school", "")).strip()
        if not degree and not school:
            continue
        education.append({
            "degree": degree,
            "school": school,
            "year": str(e.get("year", "")).strip(),
        })
    out["education"] = education
    return out


def _extract_json(text: str) -> dict:
    """Best-effort JSON recovery, mirroring ai_engine.score_job_fit fallbacks."""
    text = text.strip()
    # Strip ```json fences if the model added them
    fenced = re.search(r"```(?:json)?\s*(\{.*\})\s*```", text, re.DOTALL)
    if fenced:
        text = fenced.group(1)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if match:
        try:
            return json.loads(match.group())
        except json.JSONDecodeError:
            pass
    raise ValueError("Model did not return parseable JSON")


def _is_heading(line: str) -> bool:
    t = line.strip().rstrip(":").strip()
    if not t or len(t) > 40:
        return False
    return t.lower() in _HEADINGS or (t.isupper() and any(c.isalpha() for c in t))


def _pack(pieces: list[str], limit: int) -> list[str]:
    out, cur = [], ""
    for piece in pieces:
        if cur and len(cur) + len(piece) > limit:
            out.append(cur)
            cur = ""
        cur += piece
    if cur:
        out.append(cur)
    return out


def _split_hard(block: str, limit: int) -> list[str]:
    """Split one oversized section on blank lines (so a role stays whole), a
    still-oversized paragraph on lines, and a single huge line on characters."""
    pieces = []
    for para in (p for p in re.split(r"(?<=\n)(?=\s*\n)", block) if p):
        if len(para) <= limit:
            pieces.append(para)
            continue
        for line in para.splitlines(keepends=True):
            pieces += [line[i:i + limit] for i in range(0, len(line), limit)]
    return _pack(pieces, limit)


def chunk_resume(text: str, limit: int = APPLE_CHUNK_CHARS) -> list[str]:
    """Split résumé text on section headings into chunks of at most `limit`
    characters, packing whole sections together where they fit."""
    sections, cur = [], []
    for line in text.splitlines(keepends=True):
        if _is_heading(line) and cur:
            sections.append("".join(cur))
            cur = []
        cur.append(line)
    if cur:
        sections.append("".join(cur))
    pieces = [p for sec in sections for p in ([sec] if len(sec) <= limit else _split_hard(sec, limit))]
    return [c.strip() for c in _pack(pieces, limit) if c.strip()]


def _key(*parts) -> tuple:
    return tuple(str(p or "").strip().lower() for p in parts)


def merge_profiles(parts: list[dict]) -> dict:
    """Merge sanitized partial profiles: the first non-empty scalar wins; lists
    are concatenated and de-duplicated (experience by title+company, merging
    bullets; education by degree+school; strings case-insensitively)."""
    out = _sanitize({})
    for p in parts:
        for f in _STR_FIELDS:
            if not out[f] and p.get(f):
                out[f] = p[f]
        for f in ("skills", "certifications"):
            seen = {s.lower() for s in out[f]}
            for s in p.get(f) or []:
                if s.lower() not in seen:
                    seen.add(s.lower())
                    out[f].append(s)
        for e in p.get("experience") or []:
            match = next((x for x in out["experience"] if _key(x["title"], x["company"]) == _key(e["title"], e["company"])), None)
            if match is None:
                out["experience"].append({**e, "bullets": list(e["bullets"])})
            else:
                match["bullets"] += [b for b in e["bullets"] if b not in match["bullets"]]
        for e in p.get("education") or []:
            if all(_key(x["degree"], x["school"]) != _key(e["degree"], e["school"]) for x in out["education"]):
                out["education"].append(dict(e))
    return out


async def parse_resume(text: str, config: dict, prompt_key: str = "resume_parse") -> dict:
    """Extract a partial profile dict from résumé-like text via the local LLM.

    `prompt_key` selects the prompt_registry template; other extractive
    sources (e.g. the LinkedIn profile importer) pass their own key. The
    template must contain a `{resume}` placeholder and request the same JSON
    schema as the "resume_parse" default.

    Returns {"profile": {...}, "warnings": [...]}. Never raises for a bad
    model response — instead returns an empty profile plus a warning so the
    user can still fill the form manually.
    """
    warnings: list[str] = []
    text = (text or "").strip()
    if not text:
        return {"profile": _sanitize({}), "warnings": ["No résumé text to parse."]}

    if ai_engine._configured_model(config, "strong") == apple_bridge.SENTINEL_MODEL:
        overhead = len(prompt_registry.render_prompt(config, prompt_key, resume=""))
        chunks = chunk_resume(text, min(APPLE_CHUNK_CHARS, APPLE_INPUT_CAP - overhead - 200))
    else:
        if len(text) > _MAX_CHARS:
            text = text[:_MAX_CHARS]
            warnings.append(
                "The text was long; only the first part was parsed. Review fields carefully."
            )
        chunks = [text]

    ai_cfg = config.get("ai", {})
    parts: list[dict] = []
    bad_json: list[str] = []
    try:
        client = await ai_engine.get_client(config, "strong")
        for chunk in chunks:
            prompt = prompt_registry.render_prompt(config, prompt_key, resume=chunk)
            response = await client.chat.completions.create(
                model=ai_engine._model(config, "strong"),
                messages=[{"role": "user", "content": prompt}],
                temperature=0.1,
                max_tokens=ai_cfg.get("max_tokens", 4096),
            )
            raw_text = (response.choices[0].message.content or "").strip()
            try:
                data = _extract_json(raw_text)
            except ValueError:
                logger.warning("Résumé parse: unparseable model output: %s", raw_text[:300])
                bad_json.append(raw_text[:120])
                continue
            parts.append(_sanitize(data if isinstance(data, dict) else {}))
    except Exception as exc:
        logger.exception("Résumé parse: LLM call failed")
        code, message = ai_engine.describe_ai_error(exc, ai_cfg.get("base_url", ""))
        detail = f"{type(exc).__name__}: {exc}"
        reason = detail if code == "error" else f"{message} ({detail})"
        return {
            "profile": _sanitize({}),
            "warnings": [f"AI extraction failed: {reason}. Fill the form manually."],
        }

    if not parts:
        return {
            "profile": _sanitize({}),
            "warnings": [
                f"The AI's reply was not valid JSON (it began: {bad_json[0]!r}). Fill the form "
                "manually or try again."
            ],
        }
    if bad_json:
        warnings.append(f"{len(bad_json)} of {len(chunks)} parts of the résumé could not be read — check the fields.")

    profile = merge_profiles(parts)
    if not profile.get("full_name") and not profile.get("experience"):
        warnings.append(
            "Little structured data was found — double-check every field below."
        )
    return {"profile": profile, "warnings": warnings}
